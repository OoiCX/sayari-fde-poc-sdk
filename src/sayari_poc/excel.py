"""Read supplier portfolios without modifying the source workbook."""

from collections.abc import Iterable
from pathlib import Path

import pycountry
from openpyxl import load_workbook

from sayari_poc.models import IngestDiagnostic, IngestResult, InputEntity

_REQUIRED_HEADERS = ("name", "address", "country")


class InputError(ValueError):
    """Safe input guidance without workbook values or private paths."""


def _text(value: object) -> str | None:
    """Normalize whitespace while preserving case and Unicode."""
    return (" ".join(value.split()) or None) if isinstance(value, str) else None


def _read_rows(
    values: Iterable[tuple[object, ...]],
    sheet: str,
    limit: int | None = None,
) -> IngestResult:
    """Retain named suppliers and account for ingestion defects.

    The limit counts retained entities, not physical rows. Workbook row numbers remain stable even
    when preceding rows are rejected.
    """
    rows = iter(values)
    headers = [_text(value) for value in next(rows, ())]
    missing = [header for header in _REQUIRED_HEADERS if header not in headers]
    if missing:
        raise InputError(f"Selected sheet: missing headers: {', '.join(missing)}")
    duplicates = [header for header in _REQUIRED_HEADERS if headers.count(header) > 1]
    if duplicates:
        raise InputError(f"Selected sheet: duplicate headers: {', '.join(duplicates)}")
    columns = {header: headers.index(header) for header in _REQUIRED_HEADERS}
    entities: list[InputEntity] = []
    exceptions: list[IngestDiagnostic] = []

    # Use Excel's own row numbers; the header is row 1, so data starts at row 2.
    numbered_rows = enumerate(rows, start=2)
    # The limit caps kept entities, so dropped malformed rows don't quietly use it up.
    while limit is None or len(entities) < limit:
        try:
            row_number, row = next(numbered_rows)
        except StopIteration:
            break
        # Treat missing trailing cells as empty, and ignore columns we don't use.
        cells = {
            header: row[index] if index < len(row) else None for header, index in columns.items()
        }
        name = _text(cells["name"])
        reasons: list[str] = []
        if name is None:
            reasons.append(
                "missing name"
                if cells["name"] is None or isinstance(cells["name"], str)
                else "invalid name: expected text"
            )
        else:
            # A bad optional field only annotates the row; the named supplier is still kept.
            address = _text(cells["address"])
            if cells["address"] is not None and not isinstance(cells["address"], str):
                reasons.append("invalid address: expected text")
            country = _text(cells["country"])
            # A country, if given, must be an uppercase ASCII alpha-3 code in the ISO registry.
            if (
                country is not None
                and (
                    len(country) != 3
                    or not country.isascii()
                    or not country.isupper()
                    or pycountry.countries.get(alpha_3=country) is None
                )
            ) or (cells["country"] is not None and not isinstance(cells["country"], str)):
                reasons.append("invalid country: expected an ISO 3166-1 alpha-3 code")
                country = None
            entities.append(
                InputEntity(
                    row_number=row_number,
                    name=name,
                    address=address,
                    country=country,
                    sheet=sheet,
                )
            )
        # One entity can carry several annotations; only dropped rows add to the row totals.
        exceptions.extend(
            IngestDiagnostic(
                row_number=str(row_number),
                sheet=sheet,
                reason=reason,
                kind="dropped" if name is None else "annotated",
            )
            for reason in reasons
        )
    return IngestResult(entities=entities, exceptions=exceptions)


def read_sheet(path: Path, sheet: str, limit: int | None = None) -> IngestResult:
    """Read one worksheet or raise safe schema guidance."""
    if limit is not None and limit < 0:
        raise InputError("Ingestion limit must be nonnegative")
    # Open the workbook read-only and read formula cells as their cached results.
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise InputError("Selected sheet not found in workbook")
        return _read_rows(workbook[sheet].iter_rows(values_only=True), sheet, limit)
    finally:
        workbook.close()


def read_all_sheets(path: Path, limit: int | None = None) -> dict[str, IngestResult]:
    """Read every worksheet as a portfolio; invalid headers fail explicitly."""
    if limit is not None and limit < 0:
        raise InputError("Ingestion limit must be nonnegative")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        results: dict[str, IngestResult] = {}
        remaining = limit
        for worksheet in workbook.worksheets:
            result = _read_rows(worksheet.iter_rows(values_only=True), worksheet.title, remaining)
            results[worksheet.title] = result
            if remaining is not None:
                remaining -= len(result.entities)
        return results
    finally:
        workbook.close()
