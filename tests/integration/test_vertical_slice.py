"""Synthetic workbook and cache exercise the real stages without authentication."""

import json
from collections.abc import Iterator
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, cast
from unittest.mock import Mock

import httpx
import pytest
from openpyxl import Workbook

from sayari_poc.cache import ResponseCache
from sayari_poc.cli import main
from sayari_poc.config import Settings
from sayari_poc.pipeline import run_pipeline
from sayari_poc.report import render_report, report_datetime
from sayari_poc.transport import AuditedTransport, BudgetExceeded, OfflineCacheMiss
from tests.ontology_support import pipeline_ontology as pipeline_ontology
from tests.p3_support import entity_payload, resolution_payload

pytestmark = pytest.mark.usefixtures("pipeline_ontology")


@pytest.fixture
def seeded_cache(
    settings: Settings,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    request: pytest.FixtureRequest,
) -> Settings:
    monkeypatch.chdir(tmp_path)
    workbook = Workbook()
    workbook.remove(cast(Any, workbook.active))
    for sheet in ["list_3", "another_a", "another_b"]:
        ws = workbook.create_sheet(sheet)
        ws.append(["name", "address", "country"])
        ws.append(["Müller 示例", "Straße 1", "DEU"])
        ws.append(["Second supplier", None, None])
    path = tmp_path / "input.xlsx"
    workbook.save(path)
    workbook.close()
    configured = settings.model_copy(
        update={
            "entity_file_path": path,
            "sayari_client_id": None,
            "sayari_client_secret": None,
            "duckdb_path": tmp_path / "analysis.duckdb",
        }
    )
    cache = ResponseCache(configured.cache_dir)
    cache.put(
        cache.key(
            httpx.Request(
                "GET",
                configured.sayari_api_base + "/v1/resolution",
                params={
                    "name": "Müller 示例",
                    "address": "Straße 1",
                    "country": "DEU",
                },
            ),
        ),
        json.dumps(
            resolution_payload(
                [
                    {
                        "entity_id": "synthetic-entity",
                        "label": "示例供应商",
                        "match_strength": {"value": "strong"},
                    }
                ]
            ),
            ensure_ascii=False,
        ).encode("utf-8"),
    )
    cache.put(
        cache.key(
            httpx.Request(
                "GET", configured.sayari_api_base + "/v1/entity_summary/synthetic-entity", params={}
            )
        ),
        json.dumps(
            entity_payload(
                {
                    "id": "synthetic-entity",
                    "label": "示例供应商",
                    "translated_label": "Example Supplier",
                    "countries": ["DEU"],
                    "psa_count": 0,
                    "degree": 7,
                    "risk": {
                        "synthetic_factor": {
                            "value": True,
                            "level": getattr(request, "param", "high"),
                            "metadata": {},
                        }
                    },
                }
            ),
            ensure_ascii=False,
        ).encode("utf-8"),
    )
    cache.put(
        cache.key(
            httpx.Request(
                "GET",
                configured.sayari_api_base + "/v1/resolution",
                params={"name": "Second supplier"},
            )
        ),
        json.dumps(resolution_payload([]), ensure_ascii=False).encode("utf-8"),
    )
    cache.put(
        cache.key(
            httpx.Request(
                "GET",
                configured.sayari_api_base + "/v1/supply_chain/upstream/synthetic-entity",
                params={
                    "max_depth": configured.max_upstream_depth,
                    "limit": configured.upstream_limit,
                },
            ),
        ),
        json.dumps(
            {
                "filters": {},
                "explored_count": 0,
                "data": {"paths": [], "entities": {}},
                "partial_results": False,
            },
            ensure_ascii=False,
        ).encode("utf-8"),
    )
    return configured


@pytest.fixture(autouse=True)
def no_auth_or_send(monkeypatch: pytest.MonkeyPatch) -> Iterator[Mock]:
    # All auth and data attempts share this transport boundary (P1).
    send = Mock(side_effect=AssertionError("Authentication and HTTP forbidden"))
    monkeypatch.setattr(AuditedTransport, "_send", send)
    yield send
    send.assert_not_called()


@pytest.mark.parametrize(
    "seeded_cache,expected_level",
    [("high", "high"), ("relevant", "relevant")],
    indirect=["seeded_cache"],
)
def test_slice_produces_findings_and_report(
    seeded_cache: Settings,
    no_auth_or_send: Mock,
    tmp_path: Path,
    expected_level: str,
) -> None:
    # A cached supplier runs through the full pipeline and produces both core artifacts.
    before = seeded_cache.entity_file_path.read_bytes()
    findings = run_pipeline(seeded_cache, limit=1, offline=True)
    assert len(cast(list[dict[str, Any]], findings.suppliers)) == 1
    supplier = cast(list[dict[str, Any]], findings.suppliers)[0]
    assert supplier["status"] == supplier["resolution_status"] == "resolved"
    assert supplier["portfolio"] == "list_3"
    assert supplier["input_address"] == "Straße 1"
    assert supplier["profile"]["max_level"] == expected_level
    data = json.loads(Path("data/processed/findings.json").read_text(encoding="utf-8"))
    assert data == findings.model_dump(mode="json")
    for key, value in {
        "shared_nodes": cast(list[object], []),
    }.items():
        assert data[key] == value
    assert data["coverage"]["list_3"]["no_data"] == 1
    assert {"generated_at", "exceptions", "manifest"} <= data.keys()
    html = Path("data/processed/report.html").read_text(encoding="utf-8")
    for text in [
        "Suppliers",
        "Müller 示例",
        "示例供应商",
        "Example Supplier",
        "synthetic_factor",
        "Suppliers requiring additional review",
        "high",
    ]:
        assert text in html
    assert "<script src" not in html and "<link" not in html
    assert str(tmp_path) not in html and "Traceback" not in html
    assert seeded_cache.entity_file_path.read_bytes() == before
    no_auth_or_send.assert_not_called()


def test_identical_cache_replay_is_byte_identical(seeded_cache: Settings) -> None:
    # Re-running on unchanged evidence keeps the original timestamp and produces identical bytes.
    assert seeded_cache.declared_generated_at is None
    started = datetime.now(UTC)
    first = run_pipeline(seeded_cache, limit=1, offline=True)
    assert started <= datetime.fromisoformat(first.generated_at) <= datetime.now(UTC)
    manifest_path = Path("data/processed/run_manifest.json")
    first_manifest = json.loads(manifest_path.read_bytes())
    paths = [
        Path("data/processed/findings.json"),
        Path("data/processed/report.html"),
        Path("data/processed/flagged_subtier_entities.csv"),
    ]
    before = [p.read_bytes() for p in paths]
    second = run_pipeline(seeded_cache, limit=1, offline=True)
    second_manifest = json.loads(manifest_path.read_bytes())
    assert first_manifest["run"]["generated_at_source"] == "observed"
    assert second_manifest["run"]["generated_at_source"] == "observed"
    assert first == second
    assert before == [p.read_bytes() for p in paths]
    html = paths[1].read_text(encoding="utf-8")
    # The replay caveat was removed from the masthead. What the assertion is really for is that the
    # two runs render the same timestamp, which the byte comparison above already proves, so only
    # the label itself is pinned here.
    assert "Report content generated:" in html
    assert "Unchanged replays retain this timestamp" not in html


def test_declared_generation_is_verbatim_and_cold_replays_match(seeded_cache: Settings) -> None:
    # Replays from a clean output folder keep the declared timestamp and match byte-for-byte.
    declared = "2026-09-19T12:38:32.773832+00:00"
    configured = seeded_cache.model_copy(update={"declared_generated_at": declared})
    output = Path("data/processed")
    names = ("findings.json", "report.html", "flagged_subtier_entities.csv", "run_manifest.json")
    first = run_pipeline(configured, limit=1, offline=True)
    before = {name: (output / name).read_bytes() for name in names}
    assert first.generated_at == declared
    assert json.loads(before["findings.json"])["generated_at"] == declared
    # The declared instant reaches the reader as a date, exactly once, and the raw machine form it
    # was declared in is not rendered anywhere.
    assert before["report.html"].count(report_datetime(declared).encode()) == 1
    assert declared.encode() not in before["report.html"]
    assert b"2026-09-19T12:38:32.773832Z" not in before["report.html"]
    for name in names:
        (output / name).unlink()
    second = run_pipeline(configured, limit=1, offline=True)
    assert first == second
    for name in names[:-1]:
        assert (output / name).read_bytes() == before[name]
    first_manifest = json.loads(before["run_manifest.json"])
    second_manifest = json.loads((output / "run_manifest.json").read_bytes())
    for manifest in (first_manifest, second_manifest):
        assert manifest["run"]["generated_at_source"] == "declared"
        del manifest["run"]["timestamp"]
    assert first_manifest == second_manifest


@pytest.mark.parametrize("previous_declaration", [None, "2026-09-18T01:02:03.123456+00:00"])
def test_declared_generation_wins_over_previous_artifact(
    seeded_cache: Settings, previous_declaration: str | None
) -> None:
    # An explicit timestamp takes precedence over reusing an earlier artifact's time.
    previous_settings = seeded_cache.model_copy(
        update={"declared_generated_at": previous_declaration}
    )
    previous = run_pipeline(previous_settings, limit=1, offline=True)
    declared = "2026-09-19T12:38:32.773832+00:00"
    configured = seeded_cache.model_copy(update={"declared_generated_at": declared})
    current = run_pipeline(configured, limit=1, offline=True)
    assert current.generated_at == declared != previous.generated_at
    assert current.model_dump(exclude={"generated_at"}) == previous.model_dump(
        exclude={"generated_at"}
    )
    output = Path("data/processed")
    assert json.loads((output / "findings.json").read_bytes())["generated_at"] == declared
    rendered = (output / "report.html").read_bytes()
    assert rendered.count(report_datetime(declared).encode()) == 1
    assert declared.encode() not in rendered
    assert (
        json.loads((output / "run_manifest.json").read_bytes())["run"]["generated_at_source"]
        == "declared"
    )


@pytest.mark.parametrize("value", ["malformed-declared-time", "2026-09-19T12:38:32.773832"])
def test_invalid_declared_generation_never_reaches_artifacts(
    seeded_cache: Settings,
    value: str,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    # An invalid declared timestamp fails before artifacts are created.
    monkeypatch.setenv("DECLARED_GENERATED_AT", value)
    pipeline = Mock(side_effect=AssertionError("Invalid configuration must fail before pipeline"))
    monkeypatch.setattr("sayari_poc.cli.run_pipeline", pipeline)
    assert main(["run", "--offline"]) == 1
    assert capsys.readouterr().err == "Configuration error: check declared_generated_at\n"
    pipeline.assert_not_called()
    assert not Path("data/processed").exists()


@pytest.mark.parametrize("stage", ["resolution", "profile"])
def test_offline_miss_raises_and_preserves_error_artifacts(
    seeded_cache: Settings,
    stage: str,
    no_auth_or_send: Mock,
) -> None:
    # A cache miss is recorded against its stage before the run raises.
    cache = ResponseCache(seeded_cache.cache_dir)
    params = {"name": "Müller 示例", "address": "Straße 1", "country": "DEU"}
    key = (
        cache.key(
            httpx.Request("GET", seeded_cache.sayari_api_base + "/v1/resolution", params=params)
        )
        if stage == "resolution"
        else cache.key(
            httpx.Request(
                "GET",
                seeded_cache.sayari_api_base + "/v1/entity_summary/synthetic-entity",
                params={},
            ),
        )
    )
    (cache.cache_dir / f"{key}.json").unlink()
    with pytest.raises(OfflineCacheMiss, match="Offline cache miss"):
        run_pipeline(seeded_cache, limit=1, offline=True)
    data = json.loads(Path("data/processed/findings.json").read_text(encoding="utf-8"))
    assert data["suppliers"][0]["status"] == "error"
    assert data["exceptions"][0]["stage"] == stage
    if stage == "profile":
        assert data["suppliers"][0]["resolution_status"] == "resolved"
    no_auth_or_send.assert_not_called()


@pytest.mark.parametrize("mode", [[], ["--refresh"], ["--offline"]])
def test_dry_run_does_not_even_construct_client(
    mode: list[str],
    seeded_cache: Settings,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    # A dry run reads local inputs but never builds a client or publishes anything.
    monkeypatch.setattr("sayari_poc.cli.load_settings", lambda: seeded_cache)
    client = Mock(side_effect=AssertionError("Dry run must not construct client"))
    monkeypatch.setattr("sayari_poc.pipeline.SayariClient", client)
    assert main(["run", "--limit", "1", "--dry-run", *mode]) == 0
    output = capsys.readouterr().out.lower()
    assert "planned calls" in output and "3" in output and "estimate" in output
    if mode:
        assert f"{mode[0]} is not executed during dry-run" in output
    client.assert_not_called()
    assert not Path("data/processed/findings.json").exists()


def test_limit_stops_ingestion_and_processing(seeded_cache: Settings) -> None:
    # The entity limit caps both the input rows kept and the processing downstream.
    from openpyxl import load_workbook

    workbook = load_workbook(seeded_cache.entity_file_path)
    workbook["list_3"].append([None, "Invalid later row", None])
    workbook.save(seeded_cache.entity_file_path)
    workbook.close()
    findings = run_pipeline(seeded_cache, limit=1, offline=True)
    assert len(cast(list[dict[str, Any]], findings.suppliers)) == 1 and findings.exceptions == []
    assert cast(dict[str, Any], findings.manifest)["input_entities"] == 1


@pytest.mark.parametrize("sheets", [None, ["another_b"], ["another_a", "another_b"]])
def test_selected_and_default_sheets(seeded_cache: Settings, sheets: list[str] | None) -> None:
    # Default and explicit sheet selections keep exactly the requested portfolios.
    findings = run_pipeline(seeded_cache, sheets=sheets, offline=True)
    selected = sheets if sheets is not None else ["list_3"]
    assert [s["portfolio"] for s in cast(list[dict[str, Any]], findings.suppliers)] == [
        s for s in selected for _ in range(2)
    ]


def test_all_sheets_and_global_limit(seeded_cache: Settings) -> None:
    # With all sheets selected, one entity limit applies across them in workbook order.
    assert (
        len(
            cast(
                list[dict[str, Any]],
                run_pipeline(seeded_cache, all_sheets=True, offline=True).suppliers,
            )
        )
        == 6
    )
    limited = run_pipeline(seeded_cache, all_sheets=True, limit=3, offline=True)
    assert [s["portfolio"] for s in cast(list[dict[str, Any]], limited.suppliers)] == [
        "another_a",
        "list_3",
        "list_3",
    ]


def test_all_sheets_preserves_malformed_header_failure(seeded_cache: Settings) -> None:
    # Selecting all sheets still fails on a malformed header before anything is published.
    from openpyxl import load_workbook

    workbook = load_workbook(seeded_cache.entity_file_path)
    workbook.create_sheet("bad_headers").append(["name"])
    workbook.save(seeded_cache.entity_file_path)
    workbook.close()
    with pytest.raises(ValueError, match="Selected sheet: missing headers"):
        run_pipeline(seeded_cache, all_sheets=True, limit=1, offline=True)
    assert not Path("data/processed/findings.json").exists()


@pytest.mark.parametrize("kind", ["weak", "no_match", "malformed", "alternate", "profile"])
def test_non_success_rows_and_diagnostics_remain_visible(seeded_cache: Settings, kind: str) -> None:
    # Weak, empty and failed outcomes stay visible in the pipeline evidence.
    cache = ResponseCache(seeded_cache.cache_dir)
    key = cache.key(
        httpx.Request(
            "GET",
            seeded_cache.sayari_api_base + "/v1/resolution",
            params={
                "name": "Müller 示例",
                "address": "Straße 1",
                "country": "DEU",
            },
        ),
    )
    payload = json.loads(cache.get(key) or b"null")
    assert payload is not None
    if kind == "weak":
        payload["data"][0]["match_strength"]["value"] = "weak"
    elif kind == "no_match":
        payload["data"] = []
    elif kind == "malformed":
        payload = {"data": [{}]}
    elif kind == "alternate":
        payload["data"].append({})
    else:
        cache.put(
            cache.key(
                httpx.Request(
                    "GET",
                    seeded_cache.sayari_api_base + "/v1/entity_summary/synthetic-entity",
                    params={},
                )
            ),
            json.dumps(entity_payload({"id": "wrong"}), ensure_ascii=False).encode("utf-8"),
        )
    cache.put(key, json.dumps(payload, ensure_ascii=False).encode("utf-8"))
    findings = run_pipeline(seeded_cache, offline=True)
    assert len(cast(list[dict[str, Any]], findings.suppliers)) == 2
    row = cast(list[dict[str, Any]], findings.suppliers)[0]
    expected = {"malformed": "error", "profile": "error", "alternate": "error"}.get(kind, kind)
    assert row["status"] == expected
    if kind == "alternate":
        # A candidate missing SDK-required fields rejects the envelope.
        assert row["error_type"] == "ValidationError"
    assert findings.exceptions
    assert "Müller 示例" in Path("data/processed/report.html").read_text(encoding="utf-8")
    if kind == "profile":
        assert row["resolution_status"] == "resolved" and row["profile"] is None


@pytest.mark.parametrize(
    "args",
    [
        ["--limit", "0"],
        ["--limit", "-1"],
        ["--offline", "--refresh"],
        ["--sheet", "list_3", "--all-sheets"],
    ],
)
def test_invalid_cli_options(args: list[str]) -> None:
    # Invalid CLI combinations fail argument parsing before execution.
    with pytest.raises(SystemExit) as error:
        main(["run", *args])
    assert error.value.code == 2


def test_repeatable_sheet_cli(
    seeded_cache: Settings,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Repeated --sheet arguments select every requested portfolio, with none lost.
    monkeypatch.setattr("sayari_poc.cli.load_settings", lambda: seeded_cache)
    assert main(["run", "--sheet", "another_b", "--sheet", "another_a", "--offline"]) == 0
    data = json.loads(Path("data/processed/findings.json").read_text(encoding="utf-8"))
    assert [s["portfolio"] for s in data["suppliers"]] == ["another_a"] * 2 + ["another_b"] * 2


def test_report_escapes_untrusted_labels(seeded_cache: Settings, tmp_path: Path) -> None:
    # Untrusted labels cannot become live HTML or references to remote assets.
    findings = run_pipeline(seeded_cache, limit=1, offline=True)
    cast(list[dict[str, Any]], findings.suppliers)[0]["input_name"] = '<script>alert("x")</script>'
    render_report(findings, tmp_path / "escaped.html")
    html = (tmp_path / "escaped.html").read_text(encoding="utf-8")
    assert '<script>alert("x")</script>' not in html and "&lt;script&gt;" in html
    assert "<script src" not in html and "<link" not in html


def test_refresh_replaces_all_cached_stages(
    seeded_cache: Settings,
    no_auth_or_send: Mock,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # A mocked refresh replaces the resolution, profile and traversal evidence.
    import httpx

    send = no_auth_or_send
    seeded_cache.sayari_client_id = "synthetic-refresh-client"
    from pydantic import SecretStr

    seeded_cache.sayari_client_secret = SecretStr("synthetic-refresh-secret")

    def respond(request: httpx.Request, *, auth: bool) -> httpx.Response:
        if auth:
            return httpx.Response(
                200,
                json={
                    "access_token": "synthetic-token-for-test-only",
                    "expires_in": 3600,
                    "token_type": "Bearer",
                },
            )
        if request.url.path == "/v1/resolution":
            payload = resolution_payload(
                [
                    {
                        "entity_id": "synthetic-entity",
                        "label": "Fresh identity",
                        "match_strength": {"value": "strong"},
                    }
                ]
            )
        elif request.url.path == "/v1/supply_chain/upstream/synthetic-entity":
            payload = {
                "filters": {},
                "explored_count": 0,
                "data": {"entities": {}, "paths": []},
                "partial_results": False,
            }
        else:
            assert request.url.path == "/v1/entity_summary/synthetic-entity"
            payload = entity_payload(
                {
                    "id": "synthetic-entity",
                    "label": "Fresh profile",
                    "countries": [],
                    "psa_count": 0,
                    "risk": {},
                }
            )
        return httpx.Response(200, json=payload, request=request)

    send.side_effect = respond
    fresh = run_pipeline(seeded_cache, limit=1, refresh=True)
    assert send.call_count == 4
    assert [call.kwargs["auth"] for call in send.call_args_list] == [True, False, False, False]
    assert cast(list[dict[str, Any]], fresh.suppliers)[0]["label"] == "Fresh identity"
    assert cast(list[dict[str, Any]], fresh.suppliers)[0]["profile"]["label"] == "Fresh profile"
    send.reset_mock()
    replay = run_pipeline(seeded_cache, limit=1, offline=True)
    assert fresh == replay
    send.assert_not_called()


def test_ingestion_diagnostics_are_not_lost(seeded_cache: Settings) -> None:
    # Dropped rows and optional-field warnings survive into the final Findings.
    from openpyxl import load_workbook

    workbook = load_workbook(seeded_cache.entity_file_path)
    ws = workbook["list_3"]
    ws.append([None, "Dropped row", None])
    ws.append(["Annotated row", None, "invalid-country"])
    workbook.save(seeded_cache.entity_file_path)
    workbook.close()
    cache = ResponseCache(seeded_cache.cache_dir)
    cache.put(
        cache.key(
            httpx.Request(
                "GET",
                seeded_cache.sayari_api_base + "/v1/resolution",
                params={"name": "Annotated row"},
            )
        ),
        json.dumps(resolution_payload([]), ensure_ascii=False).encode("utf-8"),
    )
    findings = run_pipeline(seeded_cache, offline=True)
    assert len(cast(list[dict[str, Any]], findings.suppliers)) == 3
    diagnostics = [e for e in findings.exceptions if e["stage"] == "ingestion"]
    assert [(e["row_number"], e["kind"]) for e in diagnostics] == [
        ("4", "dropped"),
        ("5", "annotated"),
    ]


def test_cli_offline_miss_exits_nonzero_without_traceback(
    seeded_cache: Settings,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    no_auth_or_send: Mock,
) -> None:
    # The CLI reports a cache miss safely and returns a failing status.
    configured = seeded_cache.model_copy(update={"cache_dir": seeded_cache.cache_dir / "empty"})
    monkeypatch.setattr("sayari_poc.cli.load_settings", lambda: configured)
    assert main(["run", "--offline", "--limit", "1"]) == 1
    captured = capsys.readouterr()
    assert "Offline cache miss" in captured.err
    assert "Traceback" not in captured.err
    no_auth_or_send.assert_not_called()


def test_cli_bad_workbook_exits_without_private_path_or_traceback(
    seeded_cache: Settings,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    # Workbook failures expose neither private paths nor tracebacks.
    seeded_cache.entity_file_path.write_bytes(b"not an Excel workbook")
    monkeypatch.setattr("sayari_poc.cli.load_settings", lambda: seeded_cache)
    assert main(["run", "--dry-run"]) == 1
    captured = capsys.readouterr()
    assert "Run failed" in captured.err and "Traceback" not in captured.err
    assert str(seeded_cache.entity_file_path) not in captured.err


def test_changed_evidence_gets_new_generation_time(seeded_cache: Settings) -> None:
    # Substantive evidence changes prevent stale timestamp reuse.
    first = run_pipeline(seeded_cache, limit=1, offline=True)
    cache = ResponseCache(seeded_cache.cache_dir)
    key = cache.key(
        httpx.Request(
            "GET", seeded_cache.sayari_api_base + "/v1/entity_summary/synthetic-entity", params={}
        )
    )
    payload = json.loads(cache.get(key) or b"null")
    assert payload is not None
    payload["label"] = "Updated evidence"
    cache.put(key, json.dumps(payload, ensure_ascii=False).encode("utf-8"))
    second = run_pipeline(seeded_cache, limit=1, offline=True)
    assert first.generated_at != second.generated_at
    assert cast(list[dict[str, Any]], first.suppliers) != cast(
        list[dict[str, Any]], second.suppliers
    )


def test_weak_row_never_inherits_a_resolved_rows_profile(seeded_cache: Settings) -> None:
    # A weak candidate cannot borrow the profile of an accepted row with the same ID.
    cache = ResponseCache(seeded_cache.cache_dir)
    cache.put(
        cache.key(
            httpx.Request(
                "GET",
                seeded_cache.sayari_api_base + "/v1/resolution",
                params={"name": "Second supplier"},
            )
        ),
        json.dumps(
            resolution_payload(
                [
                    {
                        "entity_id": "synthetic-entity",
                        "label": "Weak candidate",
                        "match_strength": {"value": "weak"},
                    }
                ]
            ),
            ensure_ascii=False,
        ).encode("utf-8"),
    )
    findings = run_pipeline(seeded_cache, offline=True)
    assert cast(list[dict[str, Any]], findings.suppliers)[0]["profile"] is not None
    assert cast(list[dict[str, Any]], findings.suppliers)[1]["resolution_status"] == "weak"
    assert cast(list[dict[str, Any]], findings.suppliers)[1]["profile"] is None
    assert cast(list[dict[str, Any]], findings.suppliers)[1]["profile_status"] == "not_requested"


def test_global_limit_does_not_read_unselected_rows() -> None:
    # The entity limit stops the reader before it consumes any unselected worksheet row.
    from sayari_poc.excel import _read_rows

    def rows() -> Iterator[tuple[str | None, str | None, str | None]]:
        yield ("name", "address", "country")
        yield ("First", None, None)
        raise AssertionError("Iterator consumed a row beyond the selected entity")

    result = _read_rows(rows(), "list_3", limit=1)
    assert len(result.entities) == 1


@pytest.mark.parametrize("countries", [[], ["DEU"], None])
def test_country_evidence_distinguishes_empty_return_from_failed_profile(
    seeded_cache: Settings,
    countries: list[str] | None,
) -> None:
    # An empty country list stays distinct from a failed profile retrieval.
    cache = ResponseCache(seeded_cache.cache_dir)
    key = cache.key(
        httpx.Request(
            "GET", seeded_cache.sayari_api_base + "/v1/entity_summary/synthetic-entity", params={}
        )
    )
    profile = json.loads(cache.get(key) or b"null")
    assert profile is not None
    if countries is None:
        del profile["countries"]
    else:
        profile["countries"] = countries
    cache.put(key, json.dumps(profile, ensure_ascii=False).encode("utf-8"))
    findings = run_pipeline(seeded_cache, limit=1, offline=True)
    html = Path("data/processed/report.html").read_text(encoding="utf-8")
    if countries is None:
        assert cast(list[dict[str, Any]], findings.suppliers)[0]["profile_status"] == "error"
        # The status enum is now a sentence; the error state must still say the risk factors are
        # unknown rather than absent.
        assert "The company profile could not be retrieved" in html
        assert "Countries: none returned" not in html
    else:
        assert cast(list[dict[str, Any]], findings.suppliers)[0]["profile_status"] == "available"
        expected = ", ".join(countries) if countries else "none returned"
        assert f"Countries: {expected}" in html


def test_budget_failure_keeps_partial_rows_and_returns_nonzero(
    seeded_cache: Settings,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Running out of budget keeps every row's outcome, and the CLI reports failure.
    original_get = AuditedTransport.handle_request

    def bounded_get(self: AuditedTransport, request: httpx.Request) -> httpx.Response:
        if request.url.path == "/v1/resolution" and request.url.params["name"] != "Second supplier":
            raise BudgetExceeded("Data call budget exhausted (1)")
        return original_get(self, request)

    monkeypatch.setattr(AuditedTransport, "handle_request", bounded_get)
    monkeypatch.setattr("sayari_poc.cli.load_settings", lambda: seeded_cache)
    assert main(["run"]) == 1
    data = json.loads(Path("data/processed/findings.json").read_text(encoding="utf-8"))
    assert [row["status"] for row in data["suppliers"]] == ["error", "no_match"]
    assert data["exceptions"][0]["error_type"] == "BudgetExceeded"
    html = Path("data/processed/report.html").read_text(encoding="utf-8")
    assert "Resolution request failed" in html
    assert "Data call budget exhausted (1)" not in html
    assert "Second supplier" in html
